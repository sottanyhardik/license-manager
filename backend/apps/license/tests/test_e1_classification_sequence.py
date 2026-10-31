"""E1 display sequence, canonical labels and matching regressions.

Display sequence (column order) and classification priority (which rule wins) are
separate. Changing the display sequence must never change a classification result.
"""
import io
import uuid
from datetime import date, timedelta
from decimal import Decimal

import openpyxl
import pytest
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import RefreshToken

from apps.core.constants import GE
from apps.core.models import CompanyModel, HeadSIONNormsModel, HSCodeModel, PurchaseStatus, SionNormClassModel
from apps.license.models import LicenseDetailsModel, LicenseExportItemModel, LicenseImportItemsModel
from apps.license.services import product_classification as pc
from apps.license.services.product_classification import (
    E5_PRODUCT_CONDITIONS,
    classify_product,
    get_classification_priority,
    pivot_category_sort_key,
)
from apps.license.views.item_pivot_report import ItemPivotReportView

User = get_user_model()
PIVOT_URL = "/api/reports/item-pivot/"

E1_SEQUENCE = [
    "OTHER CONFECTIONERY",
    "DWP/SWP",
    "WPC",
    "RBD PAL",
    "ESSENTIAL OIL",
    "APPLE DIES",
    "CITRIC ACID",
    "PP",
    "PAPER",
    "CMC",
]

# (HSN, description, norm, expected) for each E1 category, preserving existing matching.
E1_MATCHING = [
    ("0802", "confectionery nut mix", "E1", "OTHER CONFECTIONERY"),
    ("08021100", "cashew nuts", "E1", "OTHER CONFECTIONERY"),
    ("0802", "plain item", "E1", "FOOD FLAVOUR"),  # no confectionery/ingredient/nut keyword
    ("0404", "whey", "E1", "DWP/SWP"),
    ("3502", "egg white", "E1", "WPC"),
    ("3502", "egg white 0404", "E1", "DWP/SWP"),  # 0404 exclusion keeps WPC out
    ("1511", "palm olein", "E1", "RBD PAL"),
    ("", "essential oil lemon", "E1", "ESSENTIAL OIL"),
    ("0813", "dried apple", "E1", "APPLE DIES"),
    ("0813", "FOOD ADDITIVES apple", "E1", "FOOD FLAVOUR"),  # APPLE DIES excludes FOOD ADDITIVES
    ("2918", "citric acid", "E1", "CITRIC ACID"),
    ("3902", "PP granules", "E1", "PP"),
    ("48", "paper roll", "E1", "PAPER"),
    ("3912", "cmc", "E1", "CMC"),
]


@pytest.mark.parametrize("hsn, description, norm, expected", E1_MATCHING)
def test_e1_matching_is_unchanged(hsn, description, norm, expected):
    assert classify_product(hsn, description, norm) == expected


def test_apple_dies_is_the_only_apple_label():
    names = {cond["name"] for cond in E5_PRODUCT_CONDITIONS}
    assert "APPLE DIES" in names
    assert "APPLE DRIES" not in names
    assert classify_product("0813", "dried apple", "E5") == "APPLE DIES"


def test_renamed_labels_are_the_only_ones_in_the_registry():
    names = {cond["name"] for cond in E5_PRODUCT_CONDITIONS}
    assert "OTHER CONFECTIONERY" in names
    assert "RBD PAL" in names
    assert "OTHER CONFECTIONERY INGREDIENTS" not in names
    assert "RBD PALMOLEIN OIL" not in names


def test_e1_sequence_is_exact():
    ordered = sorted(E1_SEQUENCE, key=lambda name: pivot_category_sort_key(name, "E1"))
    assert ordered == E1_SEQUENCE


def test_a3627_block_precedes_the_e1_block():
    a3627 = ["RUTILE", "ALUMINIUM OXIDE", "SODA ASH", "TITANIUM DIOXIDE", "PP"]
    assert max(pivot_category_sort_key(n, "A3627") for n in a3627) < min(
        pivot_category_sort_key(n, "E1") for n in E1_SEQUENCE
    )


def test_remaining_categories_follow_cmc_in_deterministic_order():
    remaining = ["SUGAR", "DIETARY FIBRE", "CEREALS", "ALUMINIUM FOIL"]
    # Under an unconfigured norm these are the "rest": they follow every configured block.
    assert min(pivot_category_sort_key(n, "E5") for n in remaining) > pivot_category_sort_key("CMC", "E1")
    # Ties break by name, so the order is deterministic.
    assert sorted(remaining, key=lambda name: (pivot_category_sort_key(name, "E5"), name)) == [
        "DIETARY FIBRE", "ALUMINIUM FOIL", "CEREALS", "SUGAR",  # DIETARY FIBRE has sort_order 1; the rest tie at 99
    ]


def test_display_sequence_does_not_change_classification(monkeypatch):
    before = [classify_product(h, d, n) for h, d, n, _ in E1_MATCHING]
    monkeypatch.setitem(pc.NORM_DISPLAY_SEQUENCES, "E1", tuple(reversed(pc.NORM_DISPLAY_SEQUENCES["E1"])))
    after = [classify_product(h, d, n) for h, d, n, _ in E1_MATCHING]
    assert before == after


def test_classification_priority_is_rule_order_not_display_order():
    assert get_classification_priority("OTHER CONFECTIONERY") is not None
    assert get_classification_priority("NOT A CATEGORY") is None
    # Rule order is the classification priority: DWP/SWP (0404) is tested before WPC (3502).
    assert get_classification_priority("DWP/SWP") < get_classification_priority("WPC")


# --- Pivot JSON and Excel order -------------------------------------------------

E1_LINES = [
    (1, "0813", "dried apple", "APPLE DIES"),
    (2, "3902", "PP granules", "PP"),
    (3, "3912", "cmc", "CMC"),
    (4, "0802", "confectionery nut mix", "OTHER CONFECTIONERY"),
    (5, "48", "paper roll", "PAPER"),
    (6, "", "sugar", "SUGAR"),
]


@pytest.fixture
def e1_licence(db):
    company = CompanyModel.objects.create(iec=f"98{uuid.uuid4().int % 10**8:08d}", name="E1 Sequence Exporter")
    status, _ = PurchaseStatus.objects.get_or_create(code=GE, defaults={"label": "Global Exim"})
    head = HeadSIONNormsModel.objects.create(name=f"E1S{uuid.uuid4().hex[:6]}")
    e1 = SionNormClassModel.objects.create(head_norm=head, norm_class="E1", is_active=True)
    lic = LicenseDetailsModel.objects.create(
        license_number="E1-SEQ-001", license_date=date.today() - timedelta(days=1),
        license_expiry_date=date.today() + timedelta(days=30), exporter=company, purchase_status=status,
    )
    LicenseExportItemModel.objects.create(license=lic, norm_class=e1, cif_fc=Decimal("100000"))
    for serial, hsn, description, _ in E1_LINES:
        code, _ = HSCodeModel.objects.get_or_create(hs_code=hsn or "99999999")
        LicenseImportItemsModel.objects.create(
            license=lic, serial_number=serial, hs_code=code, description=description,
            quantity=Decimal("100"), available_quantity=Decimal("100"), cif_fc=Decimal("50"),
        )
    return lic


@pytest.mark.django_db
def test_pivot_json_orders_e1_categories_by_display_sequence(e1_licence):
    report = ItemPivotReportView().generate_report(min_balance=0, license_status="all")
    group = next(g for g in report["groups"] if any(l["license_number"] == "E1-SEQ-001" for l in g["licenses"]))
    categories = [item["product_category"] for item in group["item_groups"]]
    expected = [name for name in ["OTHER CONFECTIONERY", "APPLE DIES", "PP", "PAPER", "CMC", "SUGAR"]]
    assert [c for c in categories if c in expected] == expected


@pytest.mark.django_db
def test_excel_headers_follow_the_same_e1_order(e1_licence):
    response = ItemPivotReportView().export_to_excel_streaming(min_balance=0, license_status="all")
    workbook = openpyxl.load_workbook(io.BytesIO(b"".join(response.streaming_content)), data_only=True)
    for sheet in workbook.worksheets:
        # Row 1 is the notification title; row 2 carries merged product-category headers.
        headers = [sheet.cell(2, c).value for c in range(10, sheet.max_column + 1) if sheet.cell(2, c).value]
        if "APPLE DIES" in headers:
            wanted = [h for h in headers if h in {"OTHER CONFECTIONERY", "APPLE DIES", "PP", "PAPER", "CMC", "SUGAR"}]
            assert wanted == ["OTHER CONFECTIONERY", "APPLE DIES", "PP", "PAPER", "CMC", "SUGAR"]
            assert "APPLE DRIES" not in headers
            return
    raise AssertionError("E1 licence not found in the Excel export")
