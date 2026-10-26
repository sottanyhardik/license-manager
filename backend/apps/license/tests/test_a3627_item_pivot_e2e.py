"""End-to-end A3627 classification: source line -> /api/reports/item-pivot/ -> Excel workbook.

Fixture lines reproduce the real descriptions and HSN variants found in the A3627
licences (Sr. Nos. 1, 6, 11, 17, 22). In the real data a serial number is unique
per licence, so each fixture line is its own licence carrying its own Sr. No.
The same source line must resolve to the same canonical category in the registry,
the JSON pivot and the exported workbook.
"""
import io
import uuid
from datetime import date, timedelta
from decimal import Decimal

import openpyxl
import pytest
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import RefreshToken

from apps.core.constants import GE
from apps.core.models import CompanyModel, HeadSIONNormsModel, HSCodeModel, PurchaseStatus, SionNormClassModel
from apps.license.models import LicenseDetailsModel, LicenseExportItemModel, LicenseImportItemsModel
from apps.license.services.a3627_classifier import classify_a3627
from apps.license.services.product_classification import classify_product
from apps.license.views import item_pivot_report as pivot_module
from apps.license.views.item_pivot_report import ItemPivotReportView

User = get_user_model()
PIVOT_URL = "/api/reports/item-pivot/"
E5_LICENCE = "E5-LEAK-001"

# (serial, HSN as stored, description, expected canonical category or None when ambiguous)
A3627_LINES = [
    (1, "32061190", "Glass Formers namely:-(a) Silica (other than Fumed Silica)28110000 (b) Borax 28401900 (c) Phosphoric Acid 28092010 (d) Arsenic Trioxide 28112920 (e) Rutile 32061190 / 32061110 [Quantity of Borax do not exceed 10%]", "RUTILE"),
    (1, "28110000", "GLASS FORMERS NAMELY (A) SILICA (OTHER THAN FUMED SILICA) (B) BORAX / BORIC ACID (NOT EXCEEDING 0.10 PER KG OF EXPORT PRODUCT)", "RUTILE"),
    (6, "28256010", "Intermediates namely:-(a) Aluminium Oxide 28182000 (b) Zinc Oxide 28170010 (c) Zirconium Oxide 28256010", "ALUMINIUM OXIDE"),
    (6, "28170010", "Intermediates namely, Zinc Oxide 28170010", "ALUMINIUM OXIDE"),
    (6, "28182000", "INTERMEDIATES NAMELY (A) ALUMINIUM OXIDE (B) ZINC OXIDE (C) RUTILE (D) ZIRCONIUM OXIDE", None),
    (11, "28362010", "Modifiers namely:- (a) Magnesium Oxide 25199040 (b) Soda Ash 28362010 (c) Sodium Nitrate 28342990 (d) Potash 28152000 (e) Potassium Nitrate 28342100", "SODA ASH"),
    (17, "32061110", "Other Special Additives namely:-(a) Clay 25080000 (b) Selenium 28049000 (c) Nucleating Agents (Titanium Dioxide other than Rutile Grade) 32061110 (d) Barium carbonate 28366000", "TITANIUM DIOXIDE"),
    (22, "39021000", "PACKING MATERIAL :- PP", "PP"),
    (22, "39021000", "Packing Material :-PP (Qty 142100.49 Kgs) 1. 48070030 / BB-Board 2/S Glazed / 900gsm in sheet form, 2. 48115900 / PP Laminated board / 380 gsm & 580 gsm in Roll form", "PP"),
    (22, "44152000", "PACKING MATERIAL (B) WOODEN SHEET / PALLETS", "WOODEN SHEETS / PALLETS"),
]


def _licence_number(index):
    return f"A3627-E2E-{index:02d}"


@pytest.fixture
def pivot_client(db):
    user = User.objects.create_user(username="a3627-pivot", email="a3627-pivot@example.com", password="RoleP@ssw0rd123")
    user.groups.add(Group.objects.get_or_create(name="REPORT_VIEWER")[0])
    client = APIClient()
    client.credentials(HTTP_AUTHORIZATION=f"Bearer {RefreshToken.for_user(user).access_token}")
    return client


@pytest.fixture
def a3627_licences(db):
    company = CompanyModel.objects.create(iec=f"99{uuid.uuid4().int % 10**8:08d}", name="A3627 E2E Exporter")
    status, _ = PurchaseStatus.objects.get_or_create(code=GE, defaults={"label": "Global Exim"})
    head = HeadSIONNormsModel.objects.create(name=f"E2E{uuid.uuid4().hex[:6]}")
    a3627 = SionNormClassModel.objects.create(head_norm=head, norm_class="A3627", is_active=True)
    e5 = SionNormClassModel.objects.create(head_norm=head, norm_class="E5", is_active=True)

    def licence(number, norm, serial, hsn, description):
        lic = LicenseDetailsModel.objects.create(
            license_number=number, license_date=date.today() - timedelta(days=1),
            license_expiry_date=date.today() + timedelta(days=30), exporter=company, purchase_status=status,
        )
        LicenseExportItemModel.objects.create(license=lic, norm_class=norm, cif_fc=Decimal("100000"))
        code, _ = HSCodeModel.objects.get_or_create(hs_code=hsn)
        LicenseImportItemsModel.objects.create(
            license=lic, serial_number=serial, hs_code=code, description=description,
            quantity=Decimal("1000"), available_quantity=Decimal("1000"), cif_fc=Decimal("500"),
        )

    for index, (serial, hsn, description, _) in enumerate(A3627_LINES):
        licence(_licence_number(index), a3627, serial, hsn, description)
    licence(E5_LICENCE, e5, 1, "32061110", "Titanium Dioxide other than Rutile Grade")
    return True


def _pivot_row(payload, number):
    for group in payload["groups"]:
        for row in group["licenses"]:
            if row.get("license_number") == number:
                return row
    raise AssertionError(f"licence {number} missing from {PIVOT_URL}")


# --- Source -> registry ----------------------------------------------------------

@pytest.mark.parametrize("serial, hsn, description, expected", A3627_LINES)
def test_registry_resolves_each_fixture_line(serial, hsn, description, expected):
    assert classify_a3627(serial, hsn, description).category == expected


@pytest.mark.parametrize("serial, hsn, description, expected", A3627_LINES)
def test_pivot_classifier_agrees_with_registry_on_each_fixture_line(serial, hsn, description, expected):
    assert classify_product(hsn, description, "A3627") == expected


def test_e5_norm_never_receives_a3627_categories():
    assert classify_product("32061110", "Titanium Dioxide other than Rutile Grade", "E5") != "TITANIUM DIOXIDE"
    assert classify_product("32061110", "Titanium Dioxide other than Rutile Grade", "E5") != "RUTILE"


# --- JSON API --------------------------------------------------------------------

@pytest.mark.django_db
def test_item_pivot_api_returns_the_registry_category_for_each_line(pivot_client, a3627_licences):
    payload = pivot_client.get(PIVOT_URL, {"min_balance": 0, "license_status": "all"}).json()
    for index, (serial, hsn, description, expected) in enumerate(A3627_LINES):
        row = _pivot_row(payload, _licence_number(index))
        cells = list(row["items"].values())
        assert len(cells) == 1, f"licence {_licence_number(index)} should show exactly one line"
        assert cells[0]["product_category"] == expected, (serial, hsn, description)
        assert str(cells[0]["hsn_code"]) == hsn


@pytest.mark.django_db
def test_e5_licence_in_same_report_is_not_reclassified_as_a3627(pivot_client, a3627_licences):
    payload = pivot_client.get(PIVOT_URL, {"min_balance": 0, "license_status": "all"}).json()
    row = _pivot_row(payload, E5_LICENCE)
    categories = {cell.get("product_category") for cell in row["items"].values()}
    assert "TITANIUM DIOXIDE" not in categories
    assert "RUTILE" not in categories


# --- Excel workbook --------------------------------------------------------------

def _workbook_for_export():
    response = ItemPivotReportView().export_to_excel_streaming(min_balance=0, license_status="all")
    data = b"".join(response.streaming_content)
    return openpyxl.load_workbook(io.BytesIO(data), data_only=True)


def _excel_categories_for_licence(workbook, number):
    """Return {hsn: {category}} for the licence's row in whichever sheet holds it.

    The category sits in row 1 at the first column of its span; later columns of
    the same span are blank, so the last non-empty header to the left applies.
    """
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=3, max_col=2):
            if str(row[1].value) != number:
                continue
            categories = {}
            current = None
            for col in range(10, sheet.max_column + 1):
                header = sheet.cell(1, col).value
                if header:
                    current = header
                if sheet.cell(2, col).value == "HSN CODE":
                    hsn = sheet.cell(row[0].row, col).value
                    if hsn not in (None, ""):
                        categories.setdefault(str(hsn), set()).add(current)
            return categories
    raise AssertionError(f"licence {number} not found in any Excel sheet")


@pytest.mark.django_db
def test_excel_export_places_each_line_under_its_category(a3627_licences):
    workbook = _workbook_for_export()
    for index, (serial, hsn, description, expected) in enumerate(A3627_LINES):
        categories = _excel_categories_for_licence(workbook, _licence_number(index))
        assert categories.get(hsn) is not None, (serial, hsn)
        expected_label = expected if expected else "UNCLASSIFIED"
        assert categories[hsn] == {expected_label}, (serial, hsn, description)


@pytest.mark.django_db
def test_excel_export_uses_no_obsolete_labels(a3627_licences):
    workbook = _workbook_for_export()
    headers = {str(c.value) for sheet in workbook.worksheets for c in sheet[1] if c.value}
    for obsolete in ("Aluminium Oxide", "Soda Ash", "Titanium Dioxide", "Wooden Sheets / Pallets", "BB-Board"):
        assert obsolete not in headers


# --- Cache invalidation ----------------------------------------------------------

@pytest.mark.django_db
def test_changing_the_registry_version_bypasses_the_cached_pivot(pivot_client, a3627_licences, monkeypatch):
    calls = []
    original = ItemPivotReportView.generate_report

    def counting(self, *args, **kwargs):
        calls.append(1)
        return original(self, *args, **kwargs)

    monkeypatch.setattr(ItemPivotReportView, "generate_report", counting)
    params = {"min_balance": 0, "license_status": "all", "marker": uuid.uuid4().hex}

    old_version, new_version = "rules-old-" + uuid.uuid4().hex, "rules-new-" + uuid.uuid4().hex

    monkeypatch.setattr(pivot_module, "registry_version", lambda: old_version)
    pivot_client.get(PIVOT_URL, params)
    pivot_client.get(PIVOT_URL, params)
    assert len(calls) == 1  # the second identical request is a cache hit

    monkeypatch.setattr(pivot_module, "registry_version", lambda: new_version)
    pivot_client.get(PIVOT_URL, params)
    assert len(calls) == 2  # a new registry version recomputes the pivot


# --- Column order and RUTILE unit price ----------------------------------------

def test_leading_categories_follow_the_a3627_sequence():
    from apps.license.services.product_classification import pivot_category_sort_key as key

    leading = ["RUTILE", "SODA ASH", "TITANIUM DIOXIDE", "PP"]
    assert sorted(leading, key=lambda name: key(name, "A3627")) == leading
    assert max(key(name, "A3627") for name in leading) < key("DIETARY FIBRE", "E5")


@pytest.mark.django_db
def test_rutile_cell_carries_balance_cif_over_balance_qty(pivot_client, a3627_licences):
    payload = pivot_client.get(PIVOT_URL, {"min_balance": 0, "license_status": "all"}).json()
    row = _pivot_row(payload, _licence_number(0))  # RUTILE line: 1000 qty, 500 CIF
    cell = next(iter(row["items"].values()))
    assert cell["product_category"] == "RUTILE"
    assert cell["unit_price"] == "0.50"


@pytest.mark.django_db
def test_excel_rutile_block_has_unit_price_after_balance_qty(a3627_licences):
    workbook = _workbook_for_export()
    for sheet in workbook.worksheets:
        rutile_cols = [c for c in range(10, sheet.max_column + 1) if sheet.cell(1, c).value == "RUTILE"]
        if not rutile_cols:
            continue
        start = rutile_cols[0]
        labels = [sheet.cell(2, c).value for c in range(start, start + 14)]
        assert labels[9] == "BALANCE QTY"
        assert labels[10] == "UNIT PRICE"
        break
    else:
        raise AssertionError("no RUTILE block in the workbook")
